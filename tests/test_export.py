"""Тесты экспорта: Excel открывается через openpyxl и сверяется с данными, JSON содержит библиотеку и журнал (PostgreSQL)."""
import io
import json
from datetime import date, datetime

import pytest
from openpyxl import load_workbook

from backend.exporter import ALL_SHEET, SHEETS, XLSX_TYPE, genres_text, status_label
from backend.models import MediaItem, ProgressLog


def add(db, title="Тайтл", category="tv_series", **over):
    item = MediaItem(title_ru=title, category=category, **over)
    db.add(item)
    db.commit()
    return item


def get_xlsx(client):
    response = client.get("/export", params={"format": "xlsx"})
    assert response.status_code == 200, response.text
    return response, load_workbook(io.BytesIO(response.content))


def rows(ws):
    """Лист как список словарей: заголовок -> значение (пустые ячейки — None)."""
    header = [c.value for c in ws[1]]
    return [dict(zip(header, [c.value for c in row])) for row in ws.iter_rows(min_row=2)]


# --- общие свойства ---

def test_default_format_is_xlsx_with_attachment_headers(client):
    response = client.get("/export")
    assert response.status_code == 200
    assert response.headers["content-type"] == XLSX_TYPE
    assert response.headers["content-disposition"] == f'attachment; filename="episodio-{date.today().isoformat()}.xlsx"'
    assert response.content[:2] == b"PK"  # xlsx — это zip


def test_invalid_format_returns_422(client):
    assert client.get("/export", params={"format": "csv"}).status_code == 422
    assert client.get("/export", params={"format": ""}).status_code == 422


def test_export_requires_api_key_when_set(client, monkeypatch):
    monkeypatch.setenv("API_KEY", "secret")
    assert client.get("/export").status_code == 401
    assert client.get("/export", headers={"X-API-Key": "secret"}).status_code == 200


# --- Excel: листы ---

def test_empty_library_gives_one_sheet_with_headers(client):
    _, wb = get_xlsx(client)
    assert wb.sheetnames == ["Библиотека"]
    assert [c.value for c in wb["Библиотека"][1]][:3] == ["Название", "Оригинальное название", "Статус"]
    assert rows(wb["Библиотека"]) == []


def test_one_sheet_per_present_category_in_fixed_order(client, db):
    add(db, "Книга", "book")
    add(db, "Аниме", "anime")
    add(db, "Фильм", "movie")
    add(db, "Сериал", "tv_series")
    add(db, "Ещё сериал", "tv_series")
    _, wb = get_xlsx(client)
    assert wb.sheetnames == [ALL_SHEET, "Аниме", "Сериалы", "Фильмы", "Книги"]
    assert [r["Название"] for r in rows(wb["Сериалы"])] == ["Ещё сериал", "Сериал"]  # по алфавиту


def test_every_category_has_its_sheet_name(client, db):
    for category, _ in SHEETS:
        add(db, f"T-{category}", category)
    _, wb = get_xlsx(client)
    assert wb.sheetnames == [ALL_SHEET] + [name for _, name in SHEETS]
    assert wb.sheetnames[1:] == ["Аниме", "Мультсериалы", "Сериалы", "Фильмы", "ТВ-шоу", "Манга", "Книги", "Аудиокниги"]


# --- Excel: заголовки и значения ---

def test_episodic_sheet_columns_and_values(client, db):
    add(
        db, "Блудливая Калифорния", "tv_series", title_original="Californication", genres=["драма", "комедия"], country="США",
        release_year=2007, current_season=7, current_episode=5, total_episodes=84, status="watching", rating=9,
        translation="LostFilm", source_link="https://example.com/x", description="Описание", comment="Моя заметка",
        author="Автор",
    )
    _, wb = get_xlsx(client)
    ws = wb["Сериалы"]
    assert [c.value for c in ws[1]] == [
        "Название", "Оригинальное название", "Статус", "Оценка", "Сезон", "Серий просмотрено", "Всего серий", "Год выхода",
        "Страна", "Жанры", "Автор", "Озвучка / перевод", "Дата завершения", "Добавлено", "Ссылка", "Описание", "Заметка",
    ]
    row = rows(ws)[0]
    assert row["Название"] == "Блудливая Калифорния"
    assert row["Оригинальное название"] == "Californication"
    assert row["Статус"] == "Смотрю"
    assert row["Оценка"] == 9 and isinstance(row["Оценка"], int)
    assert (row["Сезон"], row["Серий просмотрено"], row["Всего серий"]) == (7, 5, 84)
    assert (row["Год выхода"], row["Страна"]) == (2007, "США")
    assert row["Жанры"] == "Драма, Комедия"
    assert row["Озвучка / перевод"] == "LostFilm"
    assert row["Ссылка"] == "https://example.com/x"
    assert (row["Описание"], row["Заметка"], row["Автор"]) == ("Описание", "Моя заметка", "Автор")


def test_reading_sheet_has_volume_and_chapters(client, db):
    add(db, "Берсерк", "manga", author="Кэнтаро Миура", current_volume=3, current_chapter=40, total_chapters=364, status="watching")
    _, wb = get_xlsx(client)
    ws = wb["Манга"]
    headers = [c.value for c in ws[1]]
    assert headers[4:7] == ["Том", "Глав прочитано", "Всего глав"]
    assert "Сезон" not in headers and "Серий просмотрено" not in headers
    row = rows(ws)[0]
    assert (row["Том"], row["Глав прочитано"], row["Всего глав"]) == (3, 40, 364)
    assert row["Статус"] == "Читаю" and row["Автор"] == "Кэнтаро Миура"


def test_movie_sheet_has_no_progress_columns(client, db):
    add(db, "Интерстеллар", "movie", status="completed", rating=10)
    _, wb = get_xlsx(client)
    headers = [c.value for c in wb["Фильмы"][1]]
    assert headers[:5] == ["Название", "Оригинальное название", "Статус", "Оценка", "Год выхода"]
    assert not {"Сезон", "Серий просмотрено", "Том", "Глав прочитано"} & set(headers)
    assert rows(wb["Фильмы"])[0]["Статус"] == "Просмотрено"


@pytest.mark.parametrize(
    "category,status,expected",
    [
        ("anime", "watching", "Смотрю"), ("anime", "completed", "Просмотрено"), ("movie", "completed", "Просмотрено"),
        ("book", "watching", "Читаю"), ("manga", "completed", "Прочитано"),
        ("audiobook", "watching", "Слушаю"), ("audiobook", "completed", "Прослушано"),
        ("anime", "planned", "В планах"), ("book", "planned", "В планах"), ("tv_show", "dropped", "Брошено"),
    ],
)
def test_status_labels(category, status, expected):
    assert status_label(category, status) == expected


def test_genres_text_is_capitalized_and_comma_separated():
    assert genres_text(["драма", "научная фантастика"]) == "Драма, Научная фантастика"
    assert genres_text([]) == "" and genres_text(None) == ""


# --- Excel: даты, пустые значения, безопасность ---

def test_dates_are_real_excel_dates_in_readable_format(client, db):
    add(db, "Т", "movie", status="completed", finished_at=datetime(2026, 3, 10, 12, 30), rating=8)
    _, wb = get_xlsx(client)
    ws = wb["Фильмы"]
    headers = [c.value for c in ws[1]]
    finished = ws.cell(row=2, column=headers.index("Дата завершения") + 1)
    added = ws.cell(row=2, column=headers.index("Добавлено") + 1)
    assert finished.value == datetime(2026, 3, 10)  # дата без времени
    assert finished.number_format == "DD.MM.YYYY"
    assert isinstance(added.value, datetime) and added.number_format == "DD.MM.YYYY"


def test_empty_values_are_empty_cells_not_zeros(client, db):
    add(db, "Пустой", "tv_series")
    _, wb = get_xlsx(client)
    row = rows(wb["Сериалы"])[0]
    for name in ("Оригинальное название", "Оценка", "Сезон", "Серий просмотрено", "Всего серий", "Год выхода", "Страна", "Автор", "Дата завершения", "Заметка"):
        assert row[name] in (None, ""), name
    assert row["Статус"] == "В планах"


@pytest.mark.parametrize("title", ["=1+1", "=HYPERLINK(\"http://evil.example\",\"x\")", "+SUM(A1)", "-2+3", "@ЯЧЕЙКА", "=cmd|' /C calc'!A0"])
def test_text_starting_with_formula_characters_stays_text(client, db, title):
    add(db, title, "tv_series", comment=title, title_original=title)
    response, wb = get_xlsx(client)
    cell = wb["Сериалы"]["A2"]
    assert cell.value == title
    assert cell.data_type == "s"  # не формула
    assert b"<f>" not in b"".join(zf for zf in _sheet_xml(response.content))  # в разметке листов формул нет


def _sheet_xml(content):
    import zipfile

    with zipfile.ZipFile(io.BytesIO(content)) as z:
        for name in z.namelist():
            if name.startswith("xl/worksheets/"):
                yield z.read(name)


def test_control_characters_and_huge_text_do_not_break_export(client, db):
    add(db, "Т\x01\x02ест", "tv_series", description="x" * 40000, comment="строка\x0bс\x07мусором")
    _, wb = get_xlsx(client)
    row = rows(wb["Сериалы"])[0]
    assert row["Название"] == "Тест"
    assert len(row["Описание"]) == 32000
    assert row["Заметка"] == "строкасмусором"


def test_large_library_exports(client, db):
    db.add_all([MediaItem(title_ru=f"Тайтл {i}", category=["anime", "movie", "book"][i % 3], rating=1 + i % 10) for i in range(600)])
    db.commit()
    _, wb = get_xlsx(client)
    assert len(rows(wb[ALL_SHEET])) == 600
    assert sum(len(rows(ws)) for ws in wb.worksheets if ws.title != ALL_SHEET) == 600


def test_first_sheet_has_every_item_with_type_and_both_progress_kinds(client, db):
    add(db, "Аниме-тайтл", "anime", current_episode=3, total_episodes=12)
    add(db, "Манга-тайтл", "manga", current_volume=2, current_chapter=14, total_chapters=40)
    add(db, "Фильм-тайтл", "movie")
    _, wb = get_xlsx(client)
    assert wb.sheetnames[0] == ALL_SHEET
    by_title = {r["Название"]: r for r in rows(wb[ALL_SHEET])}
    assert set(by_title) == {"Аниме-тайтл", "Манга-тайтл", "Фильм-тайтл"}
    assert by_title["Аниме-тайтл"]["Тип"] == "Аниме" and by_title["Аниме-тайтл"]["Серий просмотрено"] == 3
    assert by_title["Манга-тайтл"]["Тип"] == "Манга" and by_title["Манга-тайтл"]["Том"] == 2
    assert by_title["Манга-тайтл"]["Глав прочитано"] == 14


# --- JSON ---

def test_json_export_contains_items_and_log(client, db):
    item = add(db, "Т", "tv_series", comment="СЕКРЕТ ВЛАДЕЛЬЦА", genres=["драма"], current_episode=0, total_episodes=3)
    client.post(f"/items/{item.id}/progress", json={"delta": 1})
    response = client.get("/export", params={"format": "json"})
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/json")
    assert response.headers["content-disposition"] == f'attachment; filename="episodio-{date.today().isoformat()}.json"'
    body = json.loads(response.content)
    assert body["app"] == "Episodio" and body["exported_at"]
    [exported] = body["items"]
    assert exported["id"] == item.id and exported["title_ru"] == "Т" and exported["genres"] == ["драма"]
    assert exported["comment"] == "СЕКРЕТ ВЛАДЕЛЬЦА"  # личная выгрузка владельца: заметки на месте
    assert exported["current_episode"] == 1 and exported["status"] == "watching"
    [entry] = body["log"]
    assert (entry["item_id"], entry["unit"], entry["number"], entry["title_ru"]) == (item.id, "episode", 1, "Т")


def test_json_export_of_empty_library(client):
    body = client.get("/export", params={"format": "json"}).json()
    assert body["items"] == [] and body["log"] == []


def test_json_log_is_oldest_first_and_removed_with_item(client, db):
    item = add(db, "Т", "tv_series", current_episode=0)
    for _ in range(3):
        client.post(f"/items/{item.id}/progress", json={"delta": 1})
    log = client.get("/export", params={"format": "json"}).json()["log"]
    assert [e["number"] for e in log] == [1, 2, 3]
    client.delete(f"/items/{item.id}")
    assert db.query(ProgressLog).count() == 0
    assert client.get("/export", params={"format": "json"}).json()["log"] == []
