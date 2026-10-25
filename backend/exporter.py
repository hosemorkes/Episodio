"""Экспорт библиотеки: Excel (по листу на категорию, русские заголовки, настоящие даты) и JSON.

Вся библиотека целиком, включая заметки: эндпоинт закрыт ключом API, как и остальные. В отличие от публичных витрин
здесь ничего не скрывается, это личная выгрузка владельца.
"""
import io
import re
from datetime import date, datetime
from typing import Any, Literal, Optional

from fastapi import APIRouter, Depends, Query, Response
from fastapi.responses import JSONResponse
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy.orm import Session

from backend.database import get_db
from backend.models import MediaItem, ProgressLog
from backend.progress import CHAPTER_CATEGORIES, EPISODE_CATEGORIES
from backend.schemas import LogEntryResponse, MediaItemResponse

router = APIRouter()

XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
DATE_FORMAT = "DD.MM.YYYY"
CELL_LIMIT = 32000  # Excel хранит в ячейке не больше 32767 символов

# Листы в порядке показа: категория -> название листа
SHEETS = [
    ("anime", "Аниме"),
    ("cartoon", "Мультсериалы"),
    ("tv_series", "Сериалы"),
    ("movie", "Фильмы"),
    ("tv_show", "ТВ-шоу"),
    ("manga", "Манга"),
    ("book", "Книги"),
    ("audiobook", "Аудиокниги"),
]
EMPTY_SHEET = "Библиотека"

# Подписи статусов: «смотрю» у сериалов и фильмов, «читаю» у манги и книг, «слушаю» у аудиокниг
STATUS_LABELS = {
    "watching": {"watch": "Смотрю", "read": "Читаю", "listen": "Слушаю"},
    "completed": {"watch": "Просмотрено", "read": "Прочитано", "listen": "Прослушано"},
    "planned": {"watch": "В планах", "read": "В планах", "listen": "В планах"},
    "dropped": {"watch": "Брошено", "read": "Брошено", "listen": "Брошено"},
}


def verb(category: str) -> str:
    return "listen" if category == "audiobook" else "read" if category in CHAPTER_CATEGORIES else "watch"


def status_label(category: str, status: str) -> str:
    return STATUS_LABELS.get(status, {}).get(verb(category), status)


# --- ячейки ---

def clean(value: Any) -> Any:
    """Текст без управляющих символов (openpyxl на них падает) и не длиннее предела ячейки."""
    if isinstance(value, str):
        return ILLEGAL_CHARACTERS_RE.sub("", value)[:CELL_LIMIT]
    return value


def put(ws: Worksheet, row: int, col: int, value: Any, number_format: Optional[str] = None) -> None:
    """Записывает значение. Текст всегда остаётся текстом: название, начинающееся с «=», не становится формулой."""
    cell = ws.cell(row=row, column=col)
    value = clean(value)
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"
    if number_format:
        cell.number_format = number_format


def genres_text(genres: Optional[list[str]]) -> str:
    """Жанры через запятую, с заглавной буквы, как в интерфейсе."""
    return ", ".join(g[:1].upper() + g[1:] for g in (genres or []) if g)


def day(value: Optional[datetime]) -> Optional[date]:
    return value.date() if value else None


# --- столбцы листов ---

def columns(category: str) -> list[tuple[str, Any, Optional[str]]]:
    """Столбцы листа категории: (заголовок, функция значения, формат ячейки). Прогресс зависит от вида категории."""
    cols: list[tuple[str, Any, Optional[str]]] = [
        ("Название", lambda i: i.title_ru, None),
        ("Оригинальное название", lambda i: i.title_original, None),
        ("Статус", lambda i: status_label(i.category, i.status), None),
        ("Оценка", lambda i: i.rating, None),
    ]
    if category in EPISODE_CATEGORIES:
        cols += [
            ("Сезон", lambda i: i.current_season, None),
            ("Серий просмотрено", lambda i: i.current_episode, None),
            ("Всего серий", lambda i: i.total_episodes, None),
        ]
    elif category in CHAPTER_CATEGORIES:
        cols += [
            ("Том", lambda i: i.current_volume, None),
            ("Глав прочитано", lambda i: i.current_chapter, None),
            ("Всего глав", lambda i: i.total_chapters, None),
        ]
    cols += [
        ("Год выхода", lambda i: i.release_year, None),
        ("Страна", lambda i: i.country, None),
        ("Жанры", lambda i: genres_text(i.genres), None),
        ("Автор", lambda i: i.author, None),
        ("Озвучка / перевод", lambda i: i.translation, None),
        ("Дата завершения", lambda i: day(i.finished_at), DATE_FORMAT),
        ("Добавлено", lambda i: day(i.created_at), DATE_FORMAT),
        ("Ссылка", lambda i: i.source_link, None),
        ("Описание", lambda i: i.description, None),
        ("Заметка", lambda i: i.comment, None),
    ]
    return cols


WIDE = {"Описание": 60, "Заметка": 50, "Ссылка": 40, "Название": 36, "Оригинальное название": 36}
HEADER_FILL = PatternFill("solid", fgColor="0F172A")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def fill_sheet(ws: Worksheet, category: str, items: list[MediaItem]) -> None:
    cols = columns(category)
    for c, (title, _, _) in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=c, value=title)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 22

    for r, item in enumerate(sorted(items, key=lambda i: (i.title_ru.casefold(), i.id)), start=2):
        for c, (title, getter, number_format) in enumerate(cols, start=1):
            value = getter(item)
            put(ws, r, c, "" if value is None else value, number_format)
            if title in ("Описание", "Заметка"):
                ws.cell(row=r, column=c).alignment = Alignment(vertical="top", wrap_text=True)

    for c, (title, _, _) in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(c)].width = WIDE.get(title, max(12, min(len(title) + 4, 24)))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{max(1, len(items) + 1)}"


def build_workbook(items: list[MediaItem]) -> Workbook:
    """Книга Excel: лист на каждую категорию, в которой есть записи; для пустой библиотеки один лист с заголовками."""
    wb = Workbook()
    wb.remove(wb.active)
    by_category: dict[str, list[MediaItem]] = {}
    for item in items:
        by_category.setdefault(item.category, []).append(item)

    for category, name in SHEETS:
        if category in by_category:
            fill_sheet(wb.create_sheet(name), category, by_category[category])
    if not wb.sheetnames:
        fill_sheet(wb.create_sheet(EMPTY_SHEET), "anime", [])
    return wb


# --- эндпоинт ---

def stamp() -> str:
    return date.today().isoformat()


def attachment(filename: str) -> dict[str, str]:
    return {"Content-Disposition": f'attachment; filename="{filename}"'}


@router.get("/export")
def export_library(
    fmt: Literal["xlsx", "json"] = Query("xlsx", alias="format", description="xlsx — Excel, json — вся библиотека и журнал"),
    db: Session = Depends(get_db),
):
    items = db.query(MediaItem).order_by(MediaItem.id).all()

    if fmt == "json":
        log_rows = (
            db.query(
                ProgressLog.id, ProgressLog.item_id, MediaItem.title_ru, MediaItem.title_original, MediaItem.category,
                ProgressLog.unit, ProgressLog.number, ProgressLog.logged_at,
            )
            .join(MediaItem, MediaItem.id == ProgressLog.item_id)
            .order_by(ProgressLog.id)
            .all()
        )
        body = {
            "app": "Episodio",
            "exported_at": datetime.now().isoformat(timespec="seconds"),
            "items": [MediaItemResponse.model_validate(i).model_dump(mode="json") for i in items],
            "log": [LogEntryResponse(**row._asdict()).model_dump(mode="json") for row in log_rows],
        }
        return JSONResponse(body, headers=attachment(f"episodio-{stamp()}.json"))

    buffer = io.BytesIO()
    build_workbook(items).save(buffer)
    return Response(buffer.getvalue(), media_type=XLSX_TYPE, headers=attachment(f"episodio-{stamp()}.xlsx"))
